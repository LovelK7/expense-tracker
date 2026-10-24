"""Legacy-layout formatting, settlement mapping and the workbook append of export_sheet.py."""

import datetime as dt
import hashlib

import openpyxl
import pytest
from openpyxl.worksheet.table import Table, TableColumn, TableFormula

import export_sheet as ex
import firebase_common as fc

UTC = dt.timezone.utc


def expense_doc(date="2026-10-25", category="Hrana", description="Konzum", cents=4599,
                paid_by="LOVEL", split="HALF", created=None, deleted=False):
    created = created or dt.datetime(2026, 10, 25, 9, 0, tzinfo=UTC)
    return {"date": date, "category": category, "description": description, "amountCents": cents,
            "paidBy": paid_by, "split": split, "createdAt": created, "updatedAt": created,
            "createdBy": paid_by, "deleted": deleted}


def settlement_doc(date="2026-10-27", sender="MARINA", cents=53469, note="listopad", deleted=False):
    created = dt.datetime(2026, 10, 27, 18, 0, tzinfo=UTC)
    return {"date": date, "from": sender, "to": "LOVEL" if sender == "MARINA" else "MARINA",
            "amountCents": cents, "note": note, "createdAt": created, "updatedAt": created,
            "createdBy": sender, "deleted": deleted}


# --- formatting ---------------------------------------------------------------------------------

def test_dates_months_and_amounts_in_the_sheet_format():
    assert ex.format_date(dt.date(2026, 3, 5)) == "5.3.2026"
    assert ex.format_date(dt.date(2023, 12, 31)) == "31.12.2023"
    assert ex.month_start(dt.date(2026, 10, 24)) == dt.date(2026, 10, 1)
    assert ex.format_amount(123456) == "1234,56"
    assert ex.format_amount(5) == "0,05"
    assert ex.format_amount(0) == "0,00"
    assert ex.format_amount(4100) == "41,00"


def test_expense_row_uses_the_croatian_values():
    fields = ex.csv_fields(ex.expense_row("abc", expense_doc()))
    assert fields == ["25.10.2026", "Hrana", "Konzum", "45,99", "Lovel", "podjela", "1.10.2026"]
    fields = ex.csv_fields(ex.expense_row("def", expense_doc(paid_by="MARINA", split="FULL_FOR_OTHER",
                                                              description="")))
    assert fields[2:6] == ["", "45,99", "Marina", "zajam"]


def test_settlement_becomes_a_prijenos_zajam_row_paid_by_the_sender():
    r = ex.settlement_row("s1", settlement_doc())
    assert ex.csv_fields(r) == ["27.10.2026", "Prijenos", "Podmirenje: listopad", "534,69", "Marina",
                                "zajam", "1.10.2026"]
    assert ex.settlement_row("s2", settlement_doc(note="  ")).description == "Podmirenje"
    assert ex.settlement_row("s3", settlement_doc(sender="LOVEL")).paid_by == "LOVEL"


@pytest.mark.parametrize("sender", ["MARINA", "LOVEL"])
def test_settlement_row_keeps_the_balance(sender):
    """The sheet's SALDO treats the zajam row like the app treats the settlement."""
    doc = settlement_doc(sender=sender, cents=1234)
    r = ex.settlement_row("s", doc)
    as_expense = {"amountCents": r.amount_cents, "paidBy": r.paid_by, "split": r.split}
    assert fc.balance_cents([as_expense]) == fc.balance_cents([], [doc])


def test_select_rows_order_since_and_deleted():
    expenses = {
        "zzz-app": expense_doc(date="2026-10-25", created=dt.datetime(2026, 10, 25, 10, tzinfo=UTC)),
        "aaa-app": expense_doc(date="2026-10-25", created=dt.datetime(2026, 10, 25, 9, tzinfo=UTC)),
        "legacy-0002": expense_doc(date="2023-01-01"),
        "legacy-0001": expense_doc(date="2023-01-05"),  # sheet order wins over date for legacy rows
        "gone": expense_doc(date="2026-10-26", deleted=True),
    }
    settlements = {"s1": settlement_doc(date="2026-10-24"), "s-del": settlement_doc(deleted=True)}
    rows = ex.select_rows(expenses, settlements)
    assert [r.source for r in rows] == ["expenses/legacy-0001", "expenses/legacy-0002",
                                        "settlements/s1", "expenses/aaa-app", "expenses/zzz-app"]
    assert [r.is_legacy for r in rows] == [True, True, False, False, False]
    since = ex.select_rows(expenses, settlements, dt.date(2026, 10, 25))
    assert [r.source for r in since] == ["expenses/aaa-app", "expenses/zzz-app"]


def test_write_csv(tmp_path):
    path = tmp_path / "out" / "export.csv"
    rows = [ex.expense_row("a", expense_doc(description='Pizza; "Napoli"'))]
    ex.write_csv(path, rows)
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    lines = raw.decode("utf-8-sig").splitlines()
    assert lines[0] == "Datum;Kategorija;Kupnja/trošak;Cijena €;Platio;Podjela/zajam;Mjesec"
    assert lines[1] == '25.10.2026;Hrana;"Pizza; ""Napoli""";45,99;Lovel;podjela;1.10.2026'
    ex.write_csv(path, rows, delimiter=",")
    assert path.read_text(encoding="utf-8-sig").splitlines()[1].startswith('25.10.2026,Hrana,"Pizza;')


# --- workbook -----------------------------------------------------------------------------------

def test_extend_ranges():
    assert ex.extend_ranges("A2:B1074 H2:H1074", 1074, 1080) == "A2:B1080 H2:H1080"
    assert ex.extend_ranges("B1:H984 E2:E1074 J5", 1074, 1080) == "B1:H984 E2:E1080 J5"
    assert ex.extend_ranges("A2:B1074", 1074, 1074) == "A2:B1074"


def test_updated_path():
    assert ex.updated_path(fc.REPO_ROOT / "private" / "Financije_ML_2023.xlsx").name == \
        "Financije_ML_2023.updated.xlsx"


MJESEC = 'IF(Table1[[#This Row],[Datum]]="","",EOMONTH(Table1[[#This Row],[Datum]],-1)+1)'


def make_workbook(path):
    """A small POPIS like the legacy one: Now, data columns, calculated Mjesec, one blank table row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "POPIS"
    header = ["Now", "Datum", "Kategorija", "Kupnja/trošak", "Cijena €", "Platio", "Podjela/zajam", "Mjesec"]
    ws.append(header)
    ws.append(["=TODAY()", dt.datetime(2023, 1, 2), "Putovanja", "Bosna", 178.8, "Lovel", "podjela", f"={MJESEC}"])
    ws.append(["=TODAY()", dt.datetime(2023, 1, 3), "Hrana", "Spar", 67.2, "Marina", "podjela", f"={MJESEC}"])
    ws.append([None] * 7 + [f"={MJESEC}"])  # blank table row
    for r in range(2, 5):
        ws.cell(r, 2).number_format = "mm-dd-yy"
        ws.cell(r, 5).number_format = "#,##0\\ [$€-1]"
    table = Table(displayName="Table1", ref="A1:H4")
    table._initialise_columns()
    table.tableColumns[7].calculatedColumnFormula = TableFormula(attr_text=MJESEC)
    ws.add_table(table)
    wb.create_sheet("SALDO")
    wb.save(path)


def test_append_to_workbook_fills_blank_rows_then_extends_the_table(tmp_path):
    source = tmp_path / "legacy.xlsx"
    make_workbook(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    rows = [ex.expense_row("a", expense_doc(date="2026-10-25", cents=4599)),
            ex.settlement_row("s", settlement_doc(date="2026-10-27")),
            ex.expense_row("b", expense_doc(date="2026-11-02", description="", cents=6001,
                                            paid_by="MARINA", split="FULL_FOR_OTHER"))]

    info = ex.append_to_workbook(source, rows)

    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest  # original untouched
    assert info["path"] == tmp_path / "legacy.updated.xlsx"
    assert info["table_ref"] == "A1:H6"
    assert info["data_rows"] == 5
    assert info["first_new_row"] == 4
    ws = openpyxl.load_workbook(info["path"])["POPIS"]
    assert ws.tables["Table1"].ref == "A1:H6"
    assert [c.value for c in ws[2]][1:7] == [dt.datetime(2023, 1, 2), "Putovanja", "Bosna", 178.8,
                                             "Lovel", "podjela"]
    assert [c.value for c in ws[4]][1:7] == [dt.datetime(2026, 10, 25), "Hrana", "Konzum", 45.99,
                                             "Lovel", "podjela"]
    assert [c.value for c in ws[5]][2:7] == ["Prijenos", "Podmirenje: listopad", 534.69, "Marina", "zajam"]
    assert [c.value for c in ws[6]][1:7] == [dt.datetime(2026, 11, 2), "Hrana", None, 60.01,
                                             "Marina", "zajam"]
    assert ws["B6"].is_date and ws["B6"].number_format == "mm-dd-yy"
    assert ws["E6"].data_type == "n" and ws["E6"].number_format == "#,##0\\ [$€-1]"
    assert ws["A6"].value == "=TODAY()"
    assert ws["H6"].value == f"={MJESEC}"


def test_append_nothing_keeps_the_table(tmp_path):
    source = tmp_path / "legacy.xlsx"
    make_workbook(source)
    info = ex.append_to_workbook(source, [])
    assert info["table_ref"] == "A1:H4"
    assert info["data_rows"] == 2
    assert info["first_new_row"] is None
