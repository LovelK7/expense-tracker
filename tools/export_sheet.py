"""Export Firestore back into the legacy spreadsheet layout (the POPIS sheet of the old workbook).

Reads every non-deleted expenses/ and settlements/ document and writes private/export_popis.csv
with the POPIS columns, in this order and with the sheet's Croatian values:

    Datum (d.M.yyyy) ; Kategorija ; Kupnja/trošak ; Cijena € (decimal comma, 2 places) ;
    Platio (Lovel/Marina) ; Podjela/zajam (podjela/zajam) ; Mjesec (first day of month, d.M.yyyy)

The CSV uses ";" as the delimiter (the list separator that goes with the decimal comma in a
Croatian locale; --delimiter "," for RFC 4180 commas) and UTF-8 with a BOM so Excel reads the
č/ć/š/ž correctly. Rows come in sheet order: the imported legacy-* rows in their original POPIS
order, then everything recorded in the app by date and creation time.

Settlements become "Prijenos" rows paid by the sender with type "zajam", which is how the legacy
sheet records a transfer (see settlement_row).

--since YYYY-MM-DD keeps only rows dated on or after that day.

--xlsx WORKBOOK additionally appends rows to the POPIS table of a copy of the legacy workbook,
written next to it as <name>.updated.xlsx; the original is never modified. Only rows that are not
already in that workbook are appended, i.e. everything except the legacy-* documents (the import
created those from this very sheet). The appended rows fill the table's blank rows first, then
the table range (and the POPIS-sourced pivot caches and data validations that tracked it) is
extended so the pivots keep working; dates are written as real Excel dates and prices as numbers,
with the formatting of the last existing row. The copy is re-opened and checked before the script
reports success.

openpyxl limitations, accepted for the copy: it drops the chart style parts, the calculation chain
(Excel rebuilds it) and the web-extension task pane. It also drops the category drop-down, which
lives in an Excel 2010 extension; this script reads that validation from the original file and
re-adds it as a standard list validation.

Read-only on Firestore. Usage (from the repo root):
    python tools/export_sheet.py --emulator [--since 2026-10-25] [--xlsx private/Financije_ML_2023.xlsx]
    python tools/export_sheet.py [--since ...] [--xlsx ...]      # the real project
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import sys
import warnings
import zipfile
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from firebase_common import (EXPENSES, LEGACY_ID_PREFIX, REPO_ROOT, SETTLEMENTS, FirestoreTarget,
                             add_target_arguments, read_collection)

DEFAULT_OUT = REPO_ROOT / "private" / "export_popis.csv"
COLUMNS = ["Datum", "Kategorija", "Kupnja/trošak", "Cijena €", "Platio", "Podjela/zajam", "Mjesec"]
PAYER_LABEL = {"LOVEL": "Lovel", "MARINA": "Marina"}
SPLIT_LABEL = {"HALF": "podjela", "FULL_FOR_OTHER": "zajam"}
TRANSFER_CATEGORY = "Prijenos"
SHEET = "POPIS"


@dataclass(frozen=True)
class SheetRow:
    """One POPIS row. paid_by and split hold the Firestore enum names."""

    date: dt.date
    category: str
    description: str
    amount_cents: int
    paid_by: str
    split: str
    source: str = ""          # "expenses/<id>" or "settlements/<id>"
    sort_key: tuple = ()

    @property
    def is_legacy(self) -> bool:
        return self.source.startswith(f"{EXPENSES}/{LEGACY_ID_PREFIX}")


def format_date(d: dt.date) -> str:
    """The sheet's d.M.yyyy: 2026-03-05 -> "5.3.2026"."""
    return f"{d.day}.{d.month}.{d.year}"


def month_start(d: dt.date) -> dt.date:
    """The Mjesec column: the first day of the row's month."""
    return d.replace(day=1)


def format_amount(cents: int) -> str:
    """Decimal comma, two places, no thousands separator: 123456 -> "1234,56"."""
    sign = "-" if cents < 0 else ""
    whole, frac = divmod(abs(cents), 100)
    return f"{sign}{whole},{frac:02d}"


def _sort_key(doc_id: str, doc: dict, date: dt.date) -> tuple:
    if doc_id.startswith(LEGACY_ID_PREFIX):
        return (0, doc_id)  # original POPIS order (legacy-NNNN is the row number)
    created = doc.get("createdAt")
    created_key = created.timestamp() if isinstance(created, dt.datetime) else 0.0
    return (1, date.isoformat(), created_key, doc_id)


def expense_row(doc_id: str, doc: dict) -> SheetRow:
    date = dt.date.fromisoformat(doc["date"])
    return SheetRow(date, doc["category"], doc["description"], int(doc["amountCents"]),
                    doc["paidBy"], doc["split"], f"{EXPENSES}/{doc_id}", _sort_key(doc_id, doc, date))


def settlement_row(doc_id: str, doc: dict) -> SheetRow:
    """A settlement as the legacy sheet records a transfer: category "Prijenos", paid by the
    sender, type "zajam". In the SALDO formula a "zajam" row counts in full in its payer's favour,
    exactly like a settlement from that person in BalanceCalculator (from MARINA: +amount, from
    LOVEL: -amount), so the sheet's balance stays equal to the app's."""
    date = dt.date.fromisoformat(doc["date"])
    note = doc.get("note", "").strip()
    description = f"Podmirenje: {note}" if note else "Podmirenje"
    return SheetRow(date, TRANSFER_CATEGORY, description, int(doc["amountCents"]),
                    doc["from"], "FULL_FOR_OTHER", f"{SETTLEMENTS}/{doc_id}",
                    _sort_key(doc_id, doc, date))


def select_rows(expenses: dict[str, dict], settlements: dict[str, dict],
                since: dt.date | None = None) -> list[SheetRow]:
    """Non-deleted expenses and settlements as sheet rows, optionally from `since` on, in sheet order."""
    rows = [expense_row(i, d) for i, d in expenses.items() if not d.get("deleted")]
    rows += [settlement_row(i, d) for i, d in settlements.items() if not d.get("deleted")]
    if since is not None:
        rows = [r for r in rows if r.date >= since]
    return sorted(rows, key=lambda r: r.sort_key)


def csv_fields(row: SheetRow) -> list[str]:
    return [format_date(row.date), row.category, row.description, format_amount(row.amount_cents),
            PAYER_LABEL[row.paid_by], SPLIT_LABEL[row.split], format_date(month_start(row.date))]


def write_csv(path: Path, rows: list[SheetRow], delimiter: str = ";") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=delimiter, lineterminator="\r\n")
        writer.writerow(COLUMNS)
        writer.writerows(csv_fields(r) for r in rows)


# ---------------------------------------------------------------------------------------------
# Workbook copy
# ---------------------------------------------------------------------------------------------

def extend_ranges(sqref: str, old_end: int, new_end: int) -> str:
    """Space-separated A1 ranges with every range that ended on old_end extended to new_end:
    extend_ranges("A2:B1074 H2:H1074 J5", 1074, 1080) == "A2:B1080 H2:H1080 J5"."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    out = []
    for part in sqref.split():
        min_col, min_row, max_col, max_row = range_boundaries(part)
        if max_row == old_end and new_end > old_end:
            out.append(f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{new_end}")
        else:
            out.append(part)
    return " ".join(out)


def updated_path(workbook: Path) -> Path:
    """private/Financije_ML_2023.xlsx -> private/Financije_ML_2023.updated.xlsx"""
    return workbook.with_name(f"{workbook.stem}.updated{workbook.suffix}")


def _sheet_xml_path(z: zipfile.ZipFile, sheet_name: str) -> str | None:
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
    wb = ElementTree.fromstring(z.read("xl/workbook.xml"))
    rels = ElementTree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    for sheet in wb.iterfind("m:sheets/m:sheet", ns):
        if sheet.get("name") == sheet_name:
            rid = sheet.get(f"{{{ns['r']}}}id")
            for rel in rels.iterfind("rel:Relationship", ns):
                if rel.get("Id") == rid:
                    target = rel.get("Target").lstrip("/")
                    return target if target.startswith("xl/") else f"xl/{target}"
    return None


def extension_list_validations(workbook: Path, sheet_name: str) -> list[dict]:
    """List validations stored in the x14 extension (openpyxl drops these on load)."""
    ns = {"x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
          "xm": "http://schemas.microsoft.com/office/excel/2006/main"}
    with zipfile.ZipFile(workbook) as z:
        path = _sheet_xml_path(z, sheet_name)
        if path is None:
            return []
        root = ElementTree.fromstring(z.read(path))
    found = []
    for dv in root.iter(f"{{{ns['x14']}}}dataValidation"):
        formula = dv.find("x14:formula1/xm:f", ns)
        sqref = dv.find("xm:sqref", ns)
        if dv.get("type") == "list" and formula is not None and sqref is not None:
            found.append({"formula1": formula.text, "sqref": sqref.text,
                          "allow_blank": dv.get("allowBlank") == "1",
                          "showErrorMessage": dv.get("showErrorMessage") == "1",
                          "showInputMessage": dv.get("showInputMessage") == "1"})
    return found


_A1_REF = re.compile(r"(?<![A-Za-z_\[\]])\$?[A-Z]{1,3}\$?\d+")


def _load(workbook: Path):
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)  # "Data Validation extension is not supported"
        return openpyxl.load_workbook(workbook)


def _popis_table(ws):
    from openpyxl.utils.cell import range_boundaries

    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        header = [ws.cell(min_row, c).value for c in range(min_col, max_col + 1)]
        if all(name in header for name in COLUMNS[:6]):
            return table, (min_col, min_row, max_col, max_row), header
    sys.exit(f"No table with the columns {COLUMNS[:6]} on sheet {SHEET}.")


def _last_data_row(ws, date_col: int, first: int, last: int) -> int:
    """The last table row with a Datum; first - 1 when the table body is empty."""
    for r in range(last, first - 1, -1):
        if ws.cell(r, date_col).value is not None:
            return r
    return first - 1


def append_to_workbook(workbook: Path, rows: list[SheetRow]) -> dict:
    """Writes updated_path(workbook): a copy with `rows` appended to the POPIS table. Returns facts
    about the result for printing; raises SystemExit if the re-opened copy does not check out."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries
    from openpyxl.worksheet.datavalidation import DataValidation

    out = updated_path(workbook)
    if out.resolve() == workbook.resolve():
        sys.exit(f"Refusing to overwrite the original workbook {workbook}")

    wb = _load(workbook)
    ws = wb[SHEET]
    table, (min_col, header_row, max_col, old_end), header = _popis_table(ws)
    col = {name: min_col + i for i, name in enumerate(header)}
    old_last = _last_data_row(ws, col["Datum"], header_row + 1, old_end)
    template = max(old_last, header_row + 1)
    calculated = {min_col + i: tc.calculatedColumnFormula.attr_text
                  for i, tc in enumerate(table.tableColumns) if tc.calculatedColumnFormula is not None}

    values_of = {
        "Datum": lambda r: dt.datetime(r.date.year, r.date.month, r.date.day),
        "Kategorija": lambda r: r.category,
        "Kupnja/trošak": lambda r: r.description or None,  # empty cell, like the legacy rows
        "Cijena €": lambda r: r.amount_cents / 100,
        "Platio": lambda r: PAYER_LABEL[r.paid_by],
        "Podjela/zajam": lambda r: SPLIT_LABEL[r.split],
    }
    data_cols = {col[name]: fn for name, fn in values_of.items()}

    for offset, row in enumerate(rows, start=1):
        r = old_last + offset
        for c in range(min_col, max_col + 1):
            source = ws.cell(template, c)
            cell = ws.cell(r, c)
            cell._style = copy(source._style)
            if c in data_cols:
                cell.value = data_cols[c](row)
            elif c in calculated:  # e.g. Mjesec: =IF(Table1[[#This Row],[Datum]]="","",EOMONTH(...)+1)
                cell.value = f"={calculated[c]}"
            elif isinstance(source.value, str) and source.value.startswith("=") \
                    and not _A1_REF.search(source.value):
                cell.value = source.value  # row-independent formulas such as Now: =TODAY()
    new_last = old_last + len(rows)
    new_end = max(old_end, new_last)

    changes = []
    if new_end > old_end:
        ref = f"{get_column_letter(min_col)}{header_row}:{get_column_letter(max_col)}{new_end}"
        changes.append(f"table {table.displayName} {table.ref} -> {ref}")
        table.ref = ref
        if table.autoFilter is not None:
            table.autoFilter.ref = ref
        for dv in ws.data_validations.dataValidation:
            before = str(dv.sqref)
            after = extend_ranges(before, old_end, new_end)
            if after != before:
                dv.sqref = after
                changes.append(f"data validation {before} -> {after}")
        for sheet in wb.worksheets:
            for pivot in sheet._pivots:
                source = pivot.cache.cacheSource.worksheetSource
                if source is None or source.sheet != SHEET or not source.ref:
                    continue
                before = source.ref
                after = extend_ranges(before, old_end, new_end)
                if after != before:
                    source.ref = after
                    changes.append(f"pivot cache of {sheet.title}!{pivot.name} {before} -> {after}")
                pivot.cache.refreshOnLoad = True

    for spec in extension_list_validations(workbook, SHEET):
        sqref = extend_ranges(spec["sqref"], old_end, new_end)
        ws.add_data_validation(DataValidation(
            type="list", formula1=spec["formula1"], allow_blank=spec["allow_blank"],
            showErrorMessage=spec["showErrorMessage"], showInputMessage=spec["showInputMessage"],
            sqref=sqref))
        changes.append(f"restored list validation {spec['formula1']} on {sqref}")

    # Save under a temporary name; it replaces <name>.updated.xlsx only once it checks out.
    partial = out.with_name(f"{out.stem}.partial{out.suffix}")
    wb.save(partial)

    # Re-open the copy and check it against the original and the appended rows.
    original = _load(workbook)[SHEET]
    copy_ws = _load(partial)[SHEET]
    copy_table, (_, _, _, copy_end), _ = _popis_table(copy_ws)
    copy_last = _last_data_row(copy_ws, col["Datum"], header_row + 1, copy_end)
    data_rows = sum(1 for r in range(header_row + 1, copy_end + 1)
                    if copy_ws.cell(r, col["Datum"]).value is not None)
    problems = []
    if copy_last != new_last or copy_end < new_last:
        problems.append(f"table {copy_table.ref} does not end with the last data row {new_last}")
    data_names = list(values_of)
    for r in range(header_row + 1, old_last + 1):
        before = [original.cell(r, col[n]).value for n in data_names]
        after = [copy_ws.cell(r, col[n]).value for n in data_names]
        if before != after:
            problems.append(f"row {r} changed: {before} -> {after}")
            break
    for offset, row in enumerate(rows, start=1):
        r = old_last + offset
        want = [values_of[n](row) for n in data_names]
        got = [copy_ws.cell(r, col[n]).value for n in data_names]
        if got != want:
            problems.append(f"appended row {r}: {got} != {want}")
            break
    if problems:
        sys.exit(f"The workbook copy failed the check (left as {partial}): " + "; ".join(problems))
    partial.replace(out)
    return {"path": out, "table_ref": copy_table.ref, "data_rows": data_rows,
            "first_new_row": old_last + 1 if rows else None, "last_data_row": new_last,
            "changes": changes, "unchanged_rows": old_last - header_row}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="CSV to write")
    parser.add_argument("--since", type=dt.date.fromisoformat, metavar="YYYY-MM-DD",
                        help="export only rows dated on or after this day")
    parser.add_argument("--xlsx", type=Path, metavar="WORKBOOK",
                        help="also append the new rows to a copy of this legacy workbook")
    parser.add_argument("--delimiter", default=";", help='CSV delimiter (default ";")')
    add_target_arguments(parser, writes=False)
    args = parser.parse_args()
    if args.xlsx is not None and not args.xlsx.is_file():
        sys.exit(f"Workbook not found: {args.xlsx}")

    with FirestoreTarget(args, writes=False) as target:
        expenses = read_collection(target.client, EXPENSES)
        settlements = read_collection(target.client, SETTLEMENTS)

    rows = select_rows(expenses, settlements, args.since)
    write_csv(args.out, rows, args.delimiter)
    n_settlements = sum(1 for r in rows if r.source.startswith(SETTLEMENTS))
    print(f"Read {len(expenses)} expenses and {len(settlements)} settlements from {target.label}")
    print(f"Wrote {len(rows)} rows ({len(rows) - n_settlements} expenses, {n_settlements} settlements"
          f"{f', dated {args.since} or later' if args.since else ''}) to {args.out}")

    if args.xlsx is not None:
        new_rows = [r for r in rows if not r.is_legacy]
        info = append_to_workbook(args.xlsx, new_rows)
        print(f"Workbook copy: {info['path']} (original {args.xlsx} untouched)")
        where = f" from row {info['first_new_row']}" if new_rows else ""
        print(f"  skipped {len(rows) - len(new_rows)} legacy-* rows (already in the workbook), "
              f"appended {len(new_rows)} row(s){where}")
        for change in info["changes"]:
            print(f"  {change}")
        print(f"  checked: re-opened, POPIS table {info['table_ref']} holds {info['data_rows']} data "
              f"rows ending at row {info['last_data_row']}; the {info['unchanged_rows']} original rows "
              "are unchanged and the appended rows match")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
