"""Export the legacy Google Drive workbook into CSV for the app's import.

Reads  private/Financije_ML_2023.xlsx  (sheets POPIS and Kategorije) and writes
  private/legacy_popis.csv        date,category,description,amount_cents,paid_by,split
  private/legacy_categories.txt   one category per line, as listed in the Kategorije sheet

Amounts are rounded to cents half-up per row. Categories are written exactly as in the sheet;
normalisation happens in the app (:domain CategoryNormalizer).

Usage (from the repo root):  python tools/export_legacy.py [--workbook PATH] [--out-dir DIR]
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import sys
import warnings
from decimal import ROUND_HALF_EVEN, ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
HEADER = ["date", "category", "description", "amount_cents", "paid_by", "split"]
EXPECTED_COLUMNS = ["Datum", "Kategorija", "Kupnja/trošak", "Cijena €", "Platio", "Podjela/zajam"]
PAID_BY = {"lovel": "LOVEL", "marina": "MARINA"}
SPLIT = {"podjela": "HALF", "zajam": "FULL_FOR_OTHER"}


def to_cents(value) -> int:
    # repr() gives the shortest decimal that round-trips the float, e.g. 20.4875 rather than
    # 20.48749999..., so half-up behaves the way a human reading the sheet would expect.
    return int((Decimal(repr(value)) * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def to_date(value) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    raise ValueError(f"not a date: {value!r}")


def read_popis(wb) -> tuple[list[list[str]], Decimal]:
    ws = wb["POPIS"]
    header = [c.value for c in ws[1]]
    try:
        col = {name: header.index(name) for name in EXPECTED_COLUMNS}
    except ValueError as e:
        sys.exit(f"POPIS header changed: {e}; got {header}")

    rows: list[list[str]] = []
    raw_balance = Decimal(0)  # what the SALDO sheet computes: raw floats, halved at the end
    for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        date, category, description, amount, payer, split = (values[col[n]] for n in EXPECTED_COLUMNS)
        if date is None and amount is None:
            continue
        if date is None or amount is None or payer is None or split is None:
            sys.exit(f"POPIS row {excel_row}: incomplete row {values[:8]}")
        try:
            paid_by = PAID_BY[str(payer).strip().lower()]
            split_type = SPLIT[str(split).strip().lower()]
        except KeyError as e:
            sys.exit(f"POPIS row {excel_row}: unknown value {e}")
        rows.append([
            to_date(date).isoformat(),
            str(category or "").strip(),
            str(description or "").strip(),
            str(to_cents(amount)),
            paid_by,
            split_type,
        ])
        sign = 1 if paid_by == "MARINA" else -1
        weight = Decimal("0.5") if split_type == "HALF" else Decimal(1)
        raw_balance += sign * weight * Decimal(repr(amount))
    return rows, raw_balance


def read_categories(wb) -> list[str]:
    ws = wb["Kategorije"]
    names = []
    for (value,) in ws.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
        if value is not None and str(value).strip():
            names.append(str(value).strip())
    return names


def cents_balance(rows: list[list[str]]) -> int:
    """Same formula as :domain BalanceCalculator (HALF difference halved once, half-even)."""
    half = full = 0
    for _, _, _, cents, paid_by, split in rows:
        signed = int(cents) if paid_by == "MARINA" else -int(cents)
        if split == "HALF":
            half += signed
        else:
            full += signed
    halved = int((Decimal(half) / 2).quantize(Decimal(1), rounding=ROUND_HALF_EVEN))
    return halved + full


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", type=Path, default=REPO_ROOT / "private" / "Financije_ML_2023.xlsx")
    parser.add_argument("--out-dir", type=Path, default=REPO_ROOT / "private")
    args = parser.parse_args()

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)  # "Data Validation extension is not supported"
        wb = openpyxl.load_workbook(args.workbook, data_only=True, read_only=False)

    rows, raw_balance = read_popis(wb)
    categories = read_categories(wb)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    csv_path = args.out_dir / "legacy_popis.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(HEADER)
        writer.writerows(rows)
    cat_path = args.out_dir / "legacy_categories.txt"
    cat_path.write_text("\n".join(categories) + "\n", encoding="utf-8")

    balance = cents_balance(rows)
    raw_cents = raw_balance * 100
    print(f"Wrote {len(rows)} rows to {csv_path}")
    print(f"Wrote {len(categories)} categories to {cat_path}")
    print(f"Dates: {min(r[0] for r in rows)} .. {max(r[0] for r in rows)}")
    print(f"Balance from cent amounts (app formula): {balance} cents  ({balance / 100:.2f} EUR)")
    print(f"Balance from raw sheet values (SALDO):   {raw_cents:.4f} cents  ({raw_balance:.6f} EUR)")
    print("Sign convention: positive = Lovel owes Marina, negative = Marina owes Lovel")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
